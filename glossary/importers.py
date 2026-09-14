"""CSV / XLSX import and export of glossary terms."""
import csv
import io
import os
import zipfile

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from django.utils.http import content_disposition_header

from .models import GlossaryTerm, GlossaryTermHistory, LegalDomain, record_history, snapshot
from .normalization import normalize

# (field, column header). Headers are bilingual; imports accept either the
# header or the field name.
COLUMNS = [
    ('term_ar', 'المصطلح بالعربية'),
    ('term_en', 'English term'),
    ('domain', 'المجال / Domain'),
    ('part_of_speech', 'نوع الكلمة / Part of speech'),
    ('definition_ar', 'التعريف بالعربية'),
    ('definition_en', 'Definition (English)'),
    ('example_ar', 'مثال بالعربية'),
    ('example_en', 'Example (English)'),
    ('synonyms_ar', 'مرادفات عربية'),
    ('synonyms_en', 'English synonyms'),
    ('forbidden_ar', 'صيغ عربية غير مقبولة'),
    ('forbidden_en', 'Forbidden English variants'),
    ('source_reference', 'المرجع القانوني / Legal reference'),
    ('notes', 'ملاحظات / Notes'),
]
FIELDS = [field for field, _ in COLUMNS]
MULTI_VALUE_FIELDS = ('synonyms_ar', 'synonyms_en', 'forbidden_ar', 'forbidden_en')
MAX_LENGTHS = {'term_ar': 500, 'term_en': 500, 'source_reference': 500}
MAX_ROWS = 5000
_FORMULA_PREFIXES = ('=', '+', '-', '@')


class ImportFileError(Exception):
    """The uploaded file can't be read as a glossary sheet."""


def _safe_cell(value):
    # Stop spreadsheet apps from executing cell contents as formulas.
    text = '' if value is None else str(value)
    return "'" + text if text.startswith(_FORMULA_PREFIXES) else text


def _unescape_cell(text):
    return text[1:] if text.startswith("'") and text[1:2] in _FORMULA_PREFIXES else text


def _split_values(text):
    parts = text.replace('\r', '\n').replace('|', '\n').split('\n')
    return [part.strip() for part in parts if part.strip()]


# ── Reading ──────────────────────────────────────────────────────────

def _raw_rows(uploaded):
    extension = os.path.splitext(uploaded.name)[1].lower()
    if extension == '.xlsx':
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(uploaded, read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise ImportFileError('تعذر قراءة ملف Excel') from exc
        sheet = workbook.worksheets[0]
        return [
            ['' if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    try:
        text = uploaded.read().decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise ImportFileError('يجب أن يكون ملف CSV بترميز UTF-8') from exc
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text))]


def read_rows(uploaded):
    """Rows as ``{field: text, '_line': n}`` dicts; the first non-empty row is the header."""
    raw = [(line, row) for line, row in enumerate(_raw_rows(uploaded), start=1) if any(row)]
    if not raw:
        raise ImportFileError('الملف فارغ')

    header_lookup = {}
    for field, label in COLUMNS:
        header_lookup[normalize(field)] = field
        header_lookup[normalize(label)] = field
    _, header = raw[0]
    positions = {
        header_lookup[normalize(cell)]: index
        for index, cell in enumerate(header) if normalize(cell) in header_lookup
    }
    missing = [label for field, label in COLUMNS[:2] if field not in positions]
    if missing:
        raise ImportFileError('الأعمدة المطلوبة غير موجودة: ' + '، '.join(missing))
    if len(raw) - 1 > MAX_ROWS:
        raise ImportFileError(f'الحد الأقصى {MAX_ROWS} صف في الملف الواحد')

    rows = []
    for line, cells in raw[1:]:
        row = {'_line': line}
        for field, index in positions.items():
            row[field] = _unescape_cell(cells[index]) if index < len(cells) else ''
        rows.append(row)
    return rows


# ── Preview & apply ──────────────────────────────────────────────────

def build_preview(rows):
    """Validate rows and flag duplicates; the result is JSON-serializable."""
    domains = {}
    for domain in LegalDomain.objects.all():
        domains[normalize(domain.name_ar)] = domain.pk
        domains[normalize(domain.name_en)] = domain.pk
    parts_of_speech = {}
    for value, label in GlossaryTerm.PartOfSpeech.choices:
        parts_of_speech[normalize(value)] = value
        parts_of_speech[normalize(label)] = value
    existing = {
        (ar, en): pk for pk, ar, en in GlossaryTerm.objects.values_list(
            'pk', 'term_ar_normalized', 'term_en_normalized')
    }

    seen = set()
    entries = []
    for row in rows:
        data = {field: row.get(field, '') for field in FIELDS}
        for field in MULTI_VALUE_FIELDS:
            data[field] = '\n'.join(_split_values(data[field]))
        errors = []

        key = (normalize(data['term_ar']), normalize(data['term_en']))
        if not key[0]:
            errors.append('المصطلح العربي مطلوب')
        if not key[1]:
            errors.append('المصطلح الإنجليزي مطلوب')
        for field, limit in MAX_LENGTHS.items():
            if len(data[field]) > limit:
                errors.append(f'النص في عمود "{dict(COLUMNS)[field]}" أطول من {limit} حرفاً')

        domain_pk = None
        if data['domain']:
            domain_pk = domains.get(normalize(data['domain']))
            if domain_pk is None:
                errors.append(f'مجال غير معروف: {data["domain"]}')
        part_of_speech = ''
        if data['part_of_speech']:
            part_of_speech = parts_of_speech.get(normalize(data['part_of_speech']), '')
            if not part_of_speech:
                errors.append(f'نوع كلمة غير معروف: {data["part_of_speech"]}')

        if all(key):
            if key in seen:
                errors.append('مكرر داخل الملف')
            seen.add(key)

        entries.append({
            'line': row['_line'],
            'data': data,
            'domain_id': domain_pk,
            'part_of_speech': part_of_speech,
            'duplicate_of': existing.get(key) if all(key) else None,
            'errors': errors,
        })
    return entries


def apply_import(entries, user, *, on_duplicate, status):
    """Create or update terms atomically. Rows with errors are skipped."""
    created = updated = skipped = 0
    with transaction.atomic():
        for entry in entries:
            if entry['errors']:
                skipped += 1
                continue
            fields = {
                **{field: entry['data'][field] for field in FIELDS if field != 'domain'},
                'domain_id': entry['domain_id'],
                'part_of_speech': entry['part_of_speech'],
            }
            existing = (
                GlossaryTerm.objects.filter(pk=entry['duplicate_of']).first()
                if entry['duplicate_of'] else None
            )
            if existing:
                if on_duplicate != 'update':
                    skipped += 1
                    continue
                before = snapshot(existing)
                for name, value in fields.items():
                    setattr(existing, name, value)
                existing.save()
                record_history(existing, GlossaryTermHistory.Action.IMPORTED, user, before)
                updated += 1
            else:
                term = GlossaryTerm(**fields, status=status, proposed_by=user)
                if status == GlossaryTerm.Status.APPROVED:
                    term.reviewed_by, term.reviewed_at = user, timezone.now()
                term.save()
                record_history(term, GlossaryTermHistory.Action.IMPORTED, user)
                created += 1
    return created, updated, skipped


# ── Export ───────────────────────────────────────────────────────────

def _export_row(term):
    values = {field: getattr(term, field) for field in FIELDS if field not in ('domain', 'part_of_speech')}
    values['domain'] = term.domain.name_ar if term.domain else ''
    values['part_of_speech'] = term.get_part_of_speech_display() if term.part_of_speech else ''
    for field in MULTI_VALUE_FIELDS:
        values[field] = ' | '.join(_split_values(values[field]))
    return [_safe_cell(values[field]) for field in FIELDS]


def export_response(terms, fmt, filename_stem):
    """An attachment with the header row plus one row per term (none for a template)."""
    header = [label for _, label in COLUMNS]
    rows = [_export_row(term) for term in terms]
    if fmt == 'xlsx':
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Glossary'
        sheet.sheet_view.rightToLeft = True
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        buffer = io.BytesIO()
        workbook.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
    else:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(header)
        writer.writerows(rows)
        # BOM so Excel opens the UTF-8 file with Arabic intact.
        response = HttpResponse('﻿' + buffer.getvalue(), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = content_disposition_header(True, f'{filename_stem}.{fmt}')
    return response
